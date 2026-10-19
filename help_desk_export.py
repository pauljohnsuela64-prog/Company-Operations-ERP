# ==========================================================
# HELP DESK - EXCEL REPORT EXPORT
# ==========================================================

from io import BytesIO

from excel_text import append_text_row

from openpyxl import Workbook

from openpyxl.styles import (
    Alignment,
    Font,
    PatternFill,
)

from openpyxl.utils import get_column_letter


# ==========================================================
# STYLES
# ==========================================================

TITLE_FILL = PatternFill(
    "solid",
    fgColor="1F4E78",
)

HEADER_FILL = PatternFill(
    "solid",
    fgColor="D9EAF7",
)


# ==========================================================
# STYLE HEADER
# ==========================================================

def _style_header(
    worksheet,
    row_number,
):

    for cell in worksheet[row_number]:

        cell.font = Font(
            bold=True,
        )

        cell.fill = HEADER_FILL

        cell.alignment = Alignment(
            vertical="center",
        )


# ==========================================================
# AUTO SIZE COLUMNS
# ==========================================================

def _auto_size_columns(worksheet):

    for column_number, column_cells in enumerate(
        worksheet.columns,
        start=1,
    ):

        max_length = 0

        # ==================================================
        # SAFE COLUMN LETTER
        #
        # This avoids the MergedCell error we encountered
        # in the IT Asset Excel export.
        # ==========================================================

        column_letter = get_column_letter(
            column_number
        )

        for cell in column_cells:

            if cell.value is None:
                continue

            value_length = len(
                str(cell.value)
            )

            max_length = max(
                max_length,
                value_length,
            )

        worksheet.column_dimensions[
            column_letter
        ].width = min(
            max_length + 2,
            40,
        )


# ==========================================================
# GENERATE HELP DESK REPORT
# ==========================================================

def generate_helpdesk_report_excel(
    summary,
    category_summary,
    technician_summary,
    tickets,
):

    workbook = Workbook()


    # ======================================================
    # SUMMARY SHEET
    # ======================================================

    summary_sheet = workbook.active

    summary_sheet.title = "Summary"

    summary_sheet.merge_cells(
        "A1:B1"
    )

    summary_sheet["A1"] = (
        "Help Desk Report"
    )

    summary_sheet["A1"].font = Font(
        bold=True,
        color="FFFFFF",
        size=14,
    )

    summary_sheet["A1"].fill = TITLE_FILL

    summary_sheet["A1"].alignment = Alignment(
        horizontal="center",
    )


    append_text_row(summary_sheet, [])

    append_text_row(summary_sheet,
        [
            "Status",
            "Ticket Count",
        ]
    )

    append_text_row(summary_sheet,
        [
            "Total Tickets",
            summary["total_tickets"],
        ]
    )

    append_text_row(summary_sheet,
        [
            "Open",
            summary["open_count"],
        ]
    )

    append_text_row(summary_sheet,
        [
            "In Progress",
            summary["in_progress_count"],
        ]
    )

    append_text_row(summary_sheet,
        [
            "Resolved",
            summary["resolved_count"],
        ]
    )

    append_text_row(summary_sheet,
        [
            "Closed",
            summary["closed_count"],
        ]
    )

    append_text_row(summary_sheet,
        [
            "Cancelled",
            summary["cancelled_count"],
        ]
    )

    _style_header(
        summary_sheet,
        3,
    )


    # ======================================================
    # BY CATEGORY SHEET
    # ======================================================

    category_sheet = workbook.create_sheet(
        "By Category"
    )

    append_text_row(category_sheet,
        [
            "Category Code",
            "Category",
            "Total",
            "Open",
            "In Progress",
            "Resolved",
            "Closed",
            "Cancelled",
        ]
    )

    for category in category_summary:

        append_text_row(category_sheet,
            [
                category[
                    "category_code"
                ],
                category[
                    "category_name"
                ],
                category[
                    "total_tickets"
                ],
                category[
                    "open_count"
                ],
                category[
                    "in_progress_count"
                ],
                category[
                    "resolved_count"
                ],
                category[
                    "closed_count"
                ],
                category[
                    "cancelled_count"
                ],
            ]
        )

    _style_header(
        category_sheet,
        1,
    )


    # ======================================================
    # BY TECHNICIAN SHEET
    # ======================================================

    technician_sheet = workbook.create_sheet(
        "By Technician"
    )

    append_text_row(technician_sheet,
        [
            "Technician",
            "Total",
            "Open",
            "In Progress",
            "Resolved",
            "Closed",
            "Cancelled",
        ]
    )

    for technician in technician_summary:

        append_text_row(technician_sheet,
            [
                technician[
                    "technician_name"
                ],
                technician[
                    "total_tickets"
                ],
                technician[
                    "open_count"
                ],
                technician[
                    "in_progress_count"
                ],
                technician[
                    "resolved_count"
                ],
                technician[
                    "closed_count"
                ],
                technician[
                    "cancelled_count"
                ],
            ]
        )

    _style_header(
        technician_sheet,
        1,
    )


    # ======================================================
    # TICKET REGISTER
    # ======================================================

    ticket_sheet = workbook.create_sheet(
        "Ticket Register"
    )

    append_text_row(ticket_sheet,
        [
            "Ticket Number",
            "Employee",
            "Employee Number",
            "Department",
            "Category",
            "Subject",
            "Description",
            "Priority",
            "Status",
            "Technician",
            "Created By",
            "Created",
            "Resolved",
            "Closed",
            "Resolution",
        ]
    )


    for ticket in tickets:

        append_text_row(ticket_sheet,
            [
                ticket[
                    "ticket_number"
                ],

                ticket[
                    "employee_name"
                ],

                ticket[
                    "employee_number"
                ],

                ticket[
                    "department_name"
                ] or "",

                ticket[
                    "category_name"
                ],

                ticket[
                    "subject"
                ],

                ticket[
                    "description"
                ],

                ticket[
                    "priority"
                ],

                ticket[
                    "status"
                ],

                ticket[
                    "assigned_to_username"
                ] or "Unassigned",

                ticket[
                    "created_by_username"
                ],

                ticket[
                    "created_at"
                ],

                ticket[
                    "resolved_at"
                ],

                ticket[
                    "closed_at"
                ],

                ticket[
                    "resolution"
                ] or "",
            ]
        )


    _style_header(
        ticket_sheet,
        1,
    )


    # ======================================================
    # DATE FORMATTING
    # ======================================================

    for row_number in range(
        2,
        ticket_sheet.max_row + 1,
    ):

        ticket_sheet.cell(
            row=row_number,
            column=12,
        ).number_format = (
            "yyyy-mm-dd hh:mm:ss"
        )

        ticket_sheet.cell(
            row=row_number,
            column=13,
        ).number_format = (
            "yyyy-mm-dd hh:mm:ss"
        )

        ticket_sheet.cell(
            row=row_number,
            column=14,
        ).number_format = (
            "yyyy-mm-dd hh:mm:ss"
        )


    # ======================================================
    # WRAP LONG TEXT
    # ======================================================

    for row_number in range(
        2,
        ticket_sheet.max_row + 1,
    ):

        ticket_sheet.cell(
            row=row_number,
            column=6,
        ).alignment = Alignment(
            vertical="top",
            wrap_text=True,
        )

        ticket_sheet.cell(
            row=row_number,
            column=7,
        ).alignment = Alignment(
            vertical="top",
            wrap_text=True,
        )

        ticket_sheet.cell(
            row=row_number,
            column=15,
        ).alignment = Alignment(
            vertical="top",
            wrap_text=True,
        )


    # ======================================================
    # FREEZE HEADERS
    # ======================================================

    category_sheet.freeze_panes = "A2"

    technician_sheet.freeze_panes = "A2"

    ticket_sheet.freeze_panes = "A2"


    # ======================================================
    # FILTERS
    # ======================================================

    category_sheet.auto_filter.ref = (
        category_sheet.dimensions
    )

    technician_sheet.auto_filter.ref = (
        technician_sheet.dimensions
    )

    ticket_sheet.auto_filter.ref = (
        ticket_sheet.dimensions
    )


    # ======================================================
    # AUTO SIZE
    # ======================================================

    for worksheet in workbook.worksheets:

        _auto_size_columns(
            worksheet
        )


    # ======================================================
    # RETURN EXCEL FILE
    # ======================================================

    output = BytesIO()

    workbook.save(
        output
    )

    output.seek(0)

    return output