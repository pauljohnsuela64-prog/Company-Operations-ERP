"""IT asset report snapshots, XLSX safety, lifecycle and reconciliation on disposable MySQL."""
from concurrent.futures import ThreadPoolExecutor
from contextlib import ExitStack
import importlib
from pathlib import Path
import re
import threading
import time
import tracemalloc
import uuid
from decimal import Decimal
from io import BytesIO
from openpyxl import load_workbook
from unittest.mock import patch

import mysql.connector
import mysql_audit_smoke as harness
from it_asset_export import generate_it_asset_report_excel
from security_helpers import stamp_session, user_state, csrf_token_for


def run(port,data_dir):
    options=dict(host='127.0.0.1',port=port,user='root',password='',connection_timeout=3)
    with mysql.connector.connect(**options) as c:
        cur=c.cursor();cur.execute('SELECT @@datadir');assert Path(cur.fetchone()[0]).resolve()==data_dir.resolve()
        cur.execute('CREATE DATABASE erp_it_reporting');cur.execute('USE erp_it_reporting')
        for name in ['core','it_assets','it_asset_assignment','it_asset_maintenance','audit_logs']:
            for statement in re.sub(r'--[^\n]*','',(harness.ROOT/'database'/f'{name}_schema.sql').read_text()).split(';'):
                if statement.strip() and not re.match(r'\s*(CREATE DATABASE|USE)\b',statement,re.I):cur.execute(statement)
        c.commit()
    options['database']='erp_it_reporting'
    def connect():return mysql.connector.connect(**options)
    def sql(query,params=(),fetch=False):
        with connect() as c:
            cur=c.cursor();cur.execute(query,params);result=cur.fetchall() if fetch else cur.lastrowid;c.commit();return result
    sql("INSERT INTO roles(role_id,role_name) VALUES(1,'Audit operator'),(2,'Viewer')")
    sql("INSERT INTO permissions(permission_code) VALUES('it_assets.view'),('it_assets.manage')")
    sql('INSERT INTO role_permissions SELECT 1,permission_id FROM permissions')
    sql("INSERT INTO role_permissions SELECT 2,permission_id FROM permissions WHERE permission_code='it_assets.view'")
    for uid in [1,2]:sql('INSERT INTO users(user_id,role_id,username,password_hash) VALUES(%s,%s,%s,\'synthetic\')',(uid,uid,f'auditor-{uid}'))
    sql("INSERT INTO employees(employee_id,employee_number,first_name,last_name) VALUES(1,'PRIVATE','private first','private last')")
    sql("INSERT INTO departments(department_id,department_code,department_name) VALUES(1,'DPT','private department')")
    sql("INSERT INTO it_asset_categories(category_id,category_code,category_name) VALUES(1,'BASE','private category')")
    with ExitStack() as stack:
        stack.enter_context(patch('dotenv.load_dotenv',return_value=False))
        stack.enter_context(patch.dict('os.environ',SECRET_KEY='isolated-it-audit',ERP_ENV='testing',ERP_DEBUG='false'))
        database=importlib.import_module('database');stack.enter_context(patch.object(database,'get_connection',side_effect=connect))
        erp=importlib.import_module('app');repo=importlib.import_module('it_asset_repository')
        stack.enter_context(patch.dict(erp.app.config,TESTING=True))
        logger=stack.enter_context(patch.object(erp.app.logger,'exception'))
        def client(uid=1):
            c=erp.app.test_client()
            with c.session_transaction() as sess:stamp_session(erp.app,sess,user_state(user_id=uid,role_id=uid,username=f'auditor-{uid}',password_hash='synthetic'))
            return c
        def post(path,fields=None,uid=1):
            c=client(uid);return c.post(path,data=dict(fields or {},csrf_token=csrf_token_for(c)))
        def logs():return sql('SELECT user_id,username,module,action,entity_type,entity_id,description FROM audit_logs ORDER BY audit_log_id',fetch=True)
        def snapshot():return [sql('SELECT * FROM '+table+' ORDER BY 1',fetch=True) for table in ['it_assets','it_asset_categories','it_asset_assignments','it_asset_maintenance','audit_logs']]
        def asset(tag, assigned=False, dept=1, category=1):
            return repo.create_it_asset(tag,tag,category,'Brand','Model',tag,dept,1 if assigned else None,
                '2026-01-01',Decimal('0.10'),'2027-01-01','ASSIGNED' if assigned else 'AVAILABLE','DAMAGED',None,1)
        def assignment(aid):return sql('SELECT assignment_id FROM it_asset_assignments WHERE asset_id=%s AND returned_at IS NULL',(aid,),True)[0][0]
        sql("INSERT INTO it_asset_categories(category_id,category_code,category_name,is_active) VALUES(2,'EMPTY','Empty inactive category',0)")
        ids=[asset('STATUS-'+state,dept=None if state=='RETIRED' else 1) for state in ['AVAILABLE','ASSIGNED','LOST','MAINTENANCE','RETIRED']]
        repo.assign_it_asset(ids[1],1,1,1)
        repo.assign_it_asset(ids[2],1,1,1);repo.mark_it_asset_lost(ids[2])
        repo.assign_it_asset(ids[3],1,1,1)
        mid=repo.create_it_asset_maintenance(ids[3],'Repair','Synthetic','Vendor','0.10','2026-10-01',None,1)
        repo.retire_it_asset(ids[4])
        def verify():
            data=repo.get_it_asset_report_data();states=[a['status'] for a in data['assets']]
            assert data['summary']['total_assets']==len(data['assets'])
            for state in ['AVAILABLE','ASSIGNED','LOST','MAINTENANCE','RETIRED']:
                assert data['summary'][state.lower()+'_count']==states.count(state)
            assert sum(row['total_assets'] for row in data['category_summary'])==len(states)
            assert sum(row['total_assets'] for row in data['department_summary'])==len(states)
            assert sum(row['total_purchase_cost'] for row in data['category_summary'])==Decimal('0.10')*len(states)
            assert not data['diagnostics'],data['diagnostics']
            wb=load_workbook(generate_it_asset_report_excel(**{k:data[k] for k in ['summary','category_summary','department_summary','assets']}))
            rows=list(wb['Asset Register'].iter_rows(min_row=2,values_only=True))
            assert [(r[0],r[8],r[9]) for r in rows]==[(r['asset_tag'],r['status'],r['condition_status']) for r in data['assets']]
            assert wb['Summary']['B4'].value==len(states)
            return data
        data=verify()
        assert all(data['summary'][state.lower()+'_count']==1 for state in ['AVAILABLE','ASSIGNED','LOST','MAINTENANCE','RETIRED'])
        assert data['category_summary'][0]['total_assets']==0 # Empty inactive category retained
        assert any(r['department_name']=='No Department' and r['total_assets']==1 for r in data['department_summary'])
        assert all(next(a for a in data['assets'] if a['asset_id']==aid)['assigned_employee_id']==1 for aid in ids[1:4])
        for kwargs,expected in [(dict(status='LOST'),[ids[2]]),(dict(category_id=2),[]),(dict(department_id=1),ids[:4]),(dict(search_text='STATUS-LOST'),[ids[2]])]:
            assert sorted(r['asset_id'] for r in repo.get_filtered_it_assets(**kwargs))==sorted(expected)
        for aid in ids:
            asset_html=client().get(f'/it/assets/{aid}').get_data(as_text=True)
            state=repo.get_it_asset_by_id(aid)
            assert (f'action="/it/assets/{aid}/assign"' in asset_html)==(state['status']=='AVAILABLE')
            assert (f'action="/it/assets/{aid}/retire"' in asset_html)==(state['status']=='AVAILABLE')
            if state['status']=='LOST':assert 'Recover Asset' in asset_html
        print('PASS five lifecycle counts, custody separate from condition/status, empty inactive category, departmentless grouping, exact Decimal totals and register/export/filter consistency')

        # Exercise committed route lifecycle and verify the report after every step.
        aid=asset('LIFECYCLE',True);token=assignment(aid);path=f'/it/assets/{aid}'
        assert len(repo.get_it_asset_assignment_history(aid))==1
        before=len(logs())
        def action(suffix,fields=None):
            assert post(path+suffix,fields).status_code==302
            verify()
        action('/lost');count=len(logs())
        action('/unassign',dict(assignment_id=token));assert len(logs())==count
        action('/restore');action('/unassign',dict(assignment_id=token))
        action('/assign',dict(employee_id='1'));new_token=assignment(aid)
        action('/unassign',dict(assignment_id=token));assert assignment(aid)==new_token
        action('/maintenance/add',dict(maintenance_type='Repair',description='Synthetic',start_date='2026-10-01',cost='0.10'))
        maintenance_id=sql('SELECT maintenance_id FROM it_asset_maintenance WHERE asset_id=%s',(aid,),True)[0][0]
        count=len(logs());action('/unassign',dict(assignment_id=new_token));assert len(logs())==count
        assert post(f'/it/assets/maintenance/{maintenance_id}/complete',dict(completed_date='2026-10-02')).status_code==302;verify()
        action('/unassign',dict(assignment_id=new_token));action('/retire');action('/restore')
        assert [r[3] for r in logs()[before:]]==['MARK_LOST','RECOVER','RETURN','ASSIGN','CREATE','COMPLETE','RETURN','RETIRE','RESTORE']
        assert all(r[0]==1 and r[2]=='IT_ASSETS' for r in logs()[before:])
        history=repo.get_it_asset_assignment_history(aid);assert len(history)==2 and all(r['returned_at'] for r in history)
        maintenance_history=repo.get_it_asset_maintenance_history(aid)
        assert len(maintenance_history)==1 and maintenance_history[0]['status']=='COMPLETED'
        print('PASS lifecycle through routes: initial history, blocked/stale returns, recovery/completion before return, retirement/restoration, histories and exactly nine success audits reflected in reports/XLSX')

        # Commit a competing write between report sections; all sections must see the old snapshot.
        original=repo.get_it_asset_report_by_category
        def concurrent_change(connection=None):
            with ThreadPoolExecutor(1) as pool:
                pool.submit(repo.mark_it_asset_lost,ids[0]).result(timeout=5)
            return original(connection=connection)
        with patch.object(repo,'get_it_asset_report_by_category',side_effect=concurrent_change):
            snapshot_data=repo.get_it_asset_report_data()
        assert snapshot_data['summary']['lost_count']==1
        assert next(a for a in snapshot_data['assets'] if a['asset_id']==ids[0])['status']=='AVAILABLE'
        assert sum(c['lost_count'] for c in snapshot_data['category_summary'])==1
        assert repo.get_it_asset_report_data()['summary']['lost_count']==2
        print('PASS real concurrent committed lifecycle write between report sections: one consistent nonblocking snapshot, fresh next report')

        # Explicit legacy corruption fixtures; diagnostics never repair or fabricate history.
        broken=asset('LEGACY-MISSING',True);sql('DELETE FROM it_asset_assignments WHERE asset_id=%s',(broken,))
        extra=asset('LEGACY-MULTIPLE',True)
        sql('INSERT INTO it_asset_assignments(asset_id,employee_id,assigned_by) VALUES(%s,1,1)',(extra,))
        bad_maint=asset('LEGACY-OPEN');sql("INSERT INTO it_asset_maintenance(asset_id,maintenance_type,description,start_date,created_by) VALUES(%s,'Repair','Synthetic','2026-10-01',1)",(bad_maint,))
        before=snapshot();findings=repo.get_it_asset_report_data()['diagnostics']
        assert {r['asset_id'] for r in findings}=={broken,extra,bad_maint}
        html=client(2).get('/it/assets/reports').get_data(as_text=True)
        assert 'History reconciliation required' in html and all(tag in html for tag in ['LEGACY-MISSING','LEGACY-MULTIPLE','LEGACY-OPEN'])
        assert f'action="/it/assets/{extra}/retire"' not in client().get(f'/it/assets/{extra}').get_data(as_text=True)
        assert snapshot()==before
        print('PASS read-only legacy missing/multiple assignment and inconsistent open maintenance diagnostics with unchanged business/audit records')

        # Actual database text -> actual endpoint -> parsed XLSX, viewer permission and fresh revocation.
        sql("UPDATE it_assets SET asset_name=%s,notes=%s WHERE asset_id=%s",('=HYPERLINK("https://invalid.example")',' \t@SUM(A1)',aid))
        response=client(2).get('/it/assets/reports/export');assert response.status_code==200
        assert 'attachment;' in response.headers['Content-Disposition']
        assert response.headers['Cache-Control']=='private, no-store'
        wb=load_workbook(BytesIO(response.data));row=next(row for row in wb['Asset Register'].iter_rows(min_row=2) if row[0].value=='LIFECYCLE')
        assert row[1].data_type=='s' and row[13].data_type=='s' and row[11].data_type=='n' and row[10].data_type=='d'
        before=snapshot()
        sql("DELETE rp FROM role_permissions rp JOIN permissions p USING(permission_id) WHERE role_id=2 AND permission_code='it_assets.view'")
        assert client(2).get('/it/assets/reports/export').location.endswith('/dashboard')
        assert client().post('/it/assets/reports/export',data={}).status_code==405
        assert snapshot()==before
        print('PASS persisted malicious text export, numeric/date types, authorized viewer download, fresh permission revocation, GET-only response and unchanged audits')

        # Small synthetic scaling observation, not a production capacity claim.
        with connect() as connection:
            cur=connection.cursor()
            cur.executemany("INSERT INTO it_assets(asset_tag,asset_name,category_id,status,condition_status,purchase_cost,created_by) VALUES(%s,%s,1,'AVAILABLE','GOOD',0.10,1)",
                            [(f'SCALE-{i}',f'Synthetic asset {i}') for i in range(1000)])
            connection.commit()
        tracemalloc.start();started=time.perf_counter()
        data=repo.get_it_asset_report_data();read_seconds=time.perf_counter()-started
        started=time.perf_counter()
        output=generate_it_asset_report_excel(**{k:data[k] for k in ['summary','category_summary','department_summary','assets']})
        export_seconds=time.perf_counter()-started
        peak=tracemalloc.get_traced_memory()[1];tracemalloc.stop()
        assert len(data['assets'])==1009 and data['summary']['total_assets']==1009
        assert load_workbook(output,read_only=True)['Asset Register'].max_row==1010
        print(f'OBSERVATION synthetic 1009 assets: snapshot {read_seconds:.3f}s, XLSX {export_seconds:.3f}s, {len(output.getvalue())} bytes, traced Python peak {peak/1024/1024:.2f} MiB; not production-scale verification')
    print('SUCCESS: IT reporting verification used disposable MySQL and synthetic records only')


if __name__=='__main__':
    harness.run=run
    harness.main()
