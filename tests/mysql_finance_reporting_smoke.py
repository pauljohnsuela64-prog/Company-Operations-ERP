"""Finance reporting/export/lifecycle on disposable MySQL with synthetic records."""
from contextlib import ExitStack
from decimal import Decimal
from io import BytesIO
import importlib
from pathlib import Path
import re
from unittest.mock import patch

import mysql.connector
from openpyxl import load_workbook
import mysql_audit_smoke as harness
from security_helpers import stamp_session, user_state, csrf_token_for
from test_finance_functional import FormValues


def verify_functional_corrections(erp, sql, c, post):
    """Real form recovery must retain links through commit, audit and payment."""
    sql("INSERT INTO suppliers(supplier_id,supplier_code,supplier_name) VALUES (1,'FORM1','Form 1'),(2,'FORM2','Form 2')")
    sql("INSERT INTO purchase_orders(purchase_order_id,po_number,supplier_id,created_by) VALUES (1,'FORM1',1,1),(2,'FORM2',2,1)")
    form = dict(expense_date='2026-12-31', category_id='2', department_id='2', supplier_id='2',
                purchase_order_id='2', payment_method='CARD', description='Form recovery', amount='0',
                reference_number='Retained reference', notes='Retained notes')
    def snapshot():
        return [sql('SELECT * FROM '+table+' ORDER BY 1', fetch=True)
                for table in ['expenses', 'department_budgets', 'audit_logs']]
    before = snapshot()
    html = post('/finance/expenses/add', form).get_data(as_text=True)
    values = FormValues(html)
    for name in ['category_id', 'department_id', 'supplier_id', 'purchase_order_id', 'payment_method']:
        assert values.selected(name) == form[name]
    assert values.values['reference_number'] == form['reference_number'] and form['notes'] in html
    assert snapshot() == before
    # A forged/stale reference stays visibly selected; it cannot turn into None.
    for name, value in [('department_id', '999'), ('supplier_id', '<invalid>'), ('purchase_order_id', '999')]:
        response = post('/finance/expenses/add', dict(form, amount='0.20', **{name: value}))
        html = response.get_data(as_text=True)
        assert response.status_code == 200 and 'invalid or no longer eligible' in html
        assert FormValues(html).selected(name) == value and snapshot() == before
    for table, field in [('departments', 'department_id'), ('suppliers', 'supplier_id'), ('expense_categories', 'category_id')]:
        sql(f'UPDATE {table} SET is_active=0 WHERE {field}=2')
    sql("UPDATE purchase_orders SET status='CANCELLED' WHERE purchase_order_id=2")
    html = post('/finance/expenses/add', dict(form, amount='0.20')).get_data(as_text=True)
    values = FormValues(html)
    for name in ['category_id', 'department_id', 'supplier_id', 'purchase_order_id']:
        assert values.selected(name) == '2'
    assert html.count('invalid or no longer eligible') == 4 and snapshot() == before
    for table, field in [('departments', 'department_id'), ('suppliers', 'supplier_id'), ('expense_categories', 'category_id')]:
        sql(f'UPDATE {table} SET is_active=1 WHERE {field}=2')
    sql("UPDATE purchase_orders SET status='DRAFT' WHERE purchase_order_id=2")
    assert post('/finance/expenses/add', dict(form, amount='0.20')).status_code == 302
    eid = sql('SELECT MAX(expense_id) FROM expenses', fetch=True)[0][0]
    assert sql('SELECT department_id,supplier_id,purchase_order_id,payment_method,amount FROM expenses WHERE expense_id=%s', (eid,), True) == [(2,2,2,'CARD',Decimal('0.20'))]
    before = snapshot()
    edited = dict(form, payment_method='BANK_TRANSFER')
    values = FormValues(post(f'/finance/expenses/{eid}/edit', edited).get_data(as_text=True))
    assert values.selected('department_id') == '2' and values.selected('payment_method') == 'BANK_TRANSFER'
    assert snapshot() == before
    assert post(f'/finance/expenses/{eid}/edit', dict(edited, amount='0.30')).status_code == 302
    assert sql('SELECT department_id,amount,payment_method FROM expenses WHERE expense_id=%s', (eid,), True) == [(2,Decimal('0.30'),'BANK_TRANSFER')]
    assert post('/finance/budgets/add', dict(department_id=2, budget_year=2026, budget_amount='1.00')).status_code == 302
    assert post(f'/finance/expenses/{eid}/approve').status_code == 302
    assert post(f'/finance/expenses/{eid}/paid').status_code == 302
    import budget_repository
    budget_id = sql('SELECT budget_id FROM department_budgets WHERE department_id=2 AND budget_year=2026', fetch=True)[0][0]
    assert budget_repository.get_department_budget_report(budget_id)[0]['spent_amount'] == Decimal('0.70')
    events = sql("SELECT action FROM audit_logs WHERE entity_type='EXPENSE' AND (entity_id=%s OR entity_id=(SELECT expense_number FROM expenses WHERE expense_id=%s)) ORDER BY audit_log_id", (str(eid), eid), True)
    assert events == [('CREATE',), ('UPDATE',), ('APPROVE',), ('MARK_PAID',)]
    for year in [2026, 2010, 2027, 1000, 9999]:
        values = FormValues(c.get(f'/finance/budgets?year={year}&department_id=2').get_data(as_text=True))
        assert values.selected('year') == str(year) and values.selected('department_id') == '2'
    response = c.get('/finance/budgets?year=NaN&department_id=2', follow_redirects=True)
    assert response.status_code==400 and 'Budget year must be between' in response.get_data(as_text=True)
    assert 'Total Budget' not in response.get_data(as_text=True)
    assert FormValues(response.get_data(as_text=True)).selected('department_id') == '2'
    # Exact exhaustion is distinct from either side; unit/Chrome checks also
    # cover amounts whose displayed percentages all round to 100.0%.
    for amount, expected in [('0.69', 'Over Budget'), ('0.70', 'Fully Used'), ('0.71', 'Warning')]:
        sql('UPDATE department_budgets SET budget_amount=%s WHERE budget_id=%s', (Decimal(amount), budget_id))
        html = c.get('/finance/budgets?year=2026&department_id=2').get_data(as_text=True)
        assert expected in html
    print('PASS real Add/Edit failed validation and correction retain references through payment/reports/audits; unavailable references, year selection and exact budget labels')


def run(port,data_dir):
    options=dict(host='127.0.0.1',port=port,user='root',password='',connection_timeout=3)
    with mysql.connector.connect(**options) as c:
        cur=c.cursor();cur.execute('SELECT @@datadir')
        assert Path(cur.fetchone()[0]).resolve()==data_dir.resolve()
        cur.execute('CREATE DATABASE erp_finance_reporting');cur.execute('USE erp_finance_reporting')
        for name in ['core','inventory','procurement','finance','budget','audit_logs']:
            for statement in re.sub(r'--[^\n]*','',(harness.ROOT/'database'/f'{name}_schema.sql').read_text()).split(';'):
                if statement.strip() and not re.match(r'\s*(CREATE DATABASE|USE)\b',statement,re.I):cur.execute(statement)
        c.commit()
    options['database']='erp_finance_reporting'
    def connect():return mysql.connector.connect(**options)
    def sql(query,params=(),fetch=False):
        with connect() as c:
            cur=c.cursor();cur.execute(query,params)
            result=cur.fetchall() if fetch else cur.lastrowid
            c.commit();return result
    sql("INSERT INTO roles(role_id,role_name) VALUES (1,'Finance Test'),(2,'Denied')")
    for uid in [1,2]:sql('INSERT INTO users(user_id,role_id,username,password_hash) VALUES (%s,%s,%s,%s)',(uid,uid,f'report-{uid}','synthetic'))
    sql("INSERT INTO permissions(permission_code) VALUES ('finance.view'),('finance.expenses.manage'),('finance.expenses.approve'),('finance.expenses.pay')")
    sql('INSERT INTO role_permissions SELECT 1,permission_id FROM permissions')
    sql("INSERT INTO departments(department_id,department_code,department_name) VALUES (1,'A','A'),(2,'B','B')")
    sql("INSERT INTO expense_categories(category_id,category_code,category_name) VALUES (1,'A','A'),(2,'B','B')")
    sql('INSERT INTO department_budgets(budget_id,department_id,budget_year,budget_amount,created_by) VALUES (1,1,2026,0.30,1)')
    with ExitStack() as stack:
        stack.enter_context(patch('dotenv.load_dotenv',return_value=False))
        stack.enter_context(patch.dict('os.environ',SECRET_KEY='isolated-finance-reporting',ERP_ENV='testing',ERP_DEBUG='false'))
        database=importlib.import_module('database')
        stack.enter_context(patch.object(database,'get_connection',side_effect=connect))
        erp=importlib.import_module('app')
        reports=importlib.import_module('finance_reporting')
        expenses=importlib.import_module('expense_repository')
        budgets=importlib.import_module('budget_repository')
        stack.enter_context(patch.dict(erp.app.config,TESTING=True))
        for summary in [expenses.get_finance_summary(),expenses.get_finance_summary_by_year(2026)]:
            assert summary['paid_count']==0 and summary['total_amount']==Decimal('0.00')
        def client(uid=1):
            c=erp.app.test_client()
            with c.session_transaction() as sess:stamp_session(erp.app,sess,user_state(user_id=uid,role_id=uid,username=f'report-{uid}',password_hash='synthetic'))
            return c
        c=client();token=csrf_token_for(c)
        def post(path,fields=None):return c.post(path,data=dict(fields or {},csrf_token=token))
        def insert(amount,status='PAID',dep=1,cat=1,when='2026-01-01'):
            return sql("INSERT INTO expenses(expense_number,expense_date,category_id,department_id,description,amount,status,paid_at,created_by) VALUES (UUID(),%s,%s,%s,'match',%s,%s,CASE WHEN %s='PAID' THEN '2030-01-01' ELSE NULL END,1)",
                       (when,cat,dep,Decimal(amount),status,status))
        first=insert('0.10');insert('0.20',cat=2,when='2026-12-31')
        insert('9','DRAFT');insert('5','APPROVED',cat=2);insert('2','CANCELLED')
        insert('100',when='2025-12-31');insert('0.40',dep=2);insert('0.50',dep=None,cat=2)
        selected=reports.filters(dict(year='2026'))
        data=reports.report(selected)
        assert data['summary']['total_amount']==Decimal('17.20') and data['summary']['total_expenses']==7
        assert data['summary']['paid_amount']==Decimal('1.20') and data['summary']['paid_count']==4
        assert sum(m['total_amount'] for m in data['monthly_expenses'])==Decimal('1.20')
        assert sum(m['total_amount'] for m in data['category_expenses'])==Decimal('1.20')
        assert expenses.get_finance_summary_by_year(2026)==data['summary']
        assert expenses.get_finance_summary()['paid_count']==5  # Dashboard is all dates.
        dashboard=importlib.import_module('dashboard_repository')
        counts=dashboard.get_dashboard_summary({'finance.view'})['finance']
        assert counts==dict(draft=1,approved=1,paid=5)
        assert c.get('/dashboard').status_code==200
        budget,records=budgets.get_department_budget_report(1)
        assert budget['spent_amount']==Decimal('0.30') and budget['remaining_amount']==Decimal('0.00')
        assert sum(r['amount'] for r in records)==budget['spent_amount']
        assert budgets.get_budget_summary(2026)['total_spent']==Decimal('0.30')
        assert reports.report(reports.filters(dict(year='2026',department_id='none')))['summary']['paid_amount']==Decimal('0.50')
        assert reports.report(reports.filters(dict(year='2026',department_id='2')))['summary']['paid_amount']==Decimal('0.40')
        query='year=2026&department_id=1&status=PAID&category_id=1&search=match'
        with patch.object(erp,'render_template',wraps=erp.render_template) as rendered:
            assert c.get('/finance/expenses?'+query).status_code==200
            assert [r['expense_id'] for r in rendered.call_args.kwargs['expenses']]==[first]
            for path in ['/finance','/finance/reports']:
                response=c.get(path+'?'+query);assert response.status_code==200
                assert rendered.call_args.kwargs['summary']['paid_amount']==Decimal('0.10')
            html=c.get('/finance/reports?'+query).get_data(as_text=True)
            for value in ['department_id=1','status=PAID','category_id=1','search=match']:assert value in html
            assert c.get('/finance/budgets?year=2026').status_code==200
            assert rendered.call_args.kwargs['budget_summary']['total_spent']==Decimal('0.30')
            assert c.get('/finance/budgets?year=2026&department_id=2').status_code==200
            assert rendered.call_args.kwargs['budget_summary']['department_count']==0
            assert c.get('/finance/budgets/1').status_code==200
            assert isinstance(rendered.call_args.kwargs['usage_percentage'],Decimal)
            old_html=' '.join(c.get('/finance?year=2025').get_data(as_text=True).split())
            assert 'Monthly Paid Expenses - 2025' in old_html
        response=c.get('/finance/reports/export?'+query)
        assert response.status_code==200 and 'attachment' in response.headers['Content-Disposition']
        assert response.headers['Cache-Control']=='private, no-store'
        wb=load_workbook(BytesIO(response.data));assert wb['Expense Records'].max_row==2
        assert 'department_id: 1' in wb['Summary']['E1'].value and 'status: PAID' in wb['Summary']['E1'].value
        assert wb['Summary']['C7'].value==0.1 and wb['Expense Records']['H2'].value==0.1
        print('PASS annual/category/status/department/search totals, all-state vs paid semantics, budget coverage, Decimal boundaries, consistent filtered UI/export')

        # Shared read-only snapshot survives a committed payment between its two queries.
        inserted=insert('0.01','APPROVED')
        sql('UPDATE department_budgets SET budget_amount=1 WHERE budget_id=1')
        original=budgets.get_department_budget_details_by_id
        def interleave(*args,**kwargs):
            result=original(*args,**kwargs)
            assert post(f'/finance/expenses/{inserted}/paid').status_code==302
            return result
        with patch.object(budgets,'get_department_budget_details_by_id',side_effect=interleave):
            old,old_rows=budgets.get_department_budget_report(1)
        assert old['spent_amount']==Decimal('0.30') and sum(r['amount'] for r in old_rows)==Decimal('0.30')
        assert budgets.get_department_budget_report(1)[0]['spent_amount']==Decimal('0.31')

        # Full real route lifecycle; reports reflect each committed state.
        form=dict(expense_date='2026-02-01',category_id='1',department_id='1',description='=1+1',amount='0.10')
        assert post('/finance/expenses/add',form).status_code==302
        eid=sql('SELECT MAX(expense_id) FROM expenses',fetch=True)[0][0]
        assert post(f'/finance/expenses/{eid}/edit',dict(form,amount='0.20')).status_code==302
        def record():return next(r for r in reports.report(selected)['expense_records'] if r['expense_id']==eid)
        assert record()['status']=='DRAFT' and record()['amount']==Decimal('0.20')
        for endpoint,status in [('approve','APPROVED'),('paid','PAID')]:
            assert post(f'/finance/expenses/{eid}/{endpoint}').status_code==302
            assert record()['status']==status
        audit_before=sql('SELECT COUNT(*) FROM audit_logs',fetch=True)
        assert post(f'/finance/expenses/{eid}/paid').status_code==302
        assert post(f'/finance/expenses/{eid}/cancel').status_code==302
        assert sql('SELECT COUNT(*) FROM audit_logs',fetch=True)==audit_before
        events=sql("SELECT action FROM audit_logs WHERE module='FINANCE' AND entity_type='EXPENSE' AND ((entity_id=%s AND action IN ('CREATE','UPDATE')) OR entity_id=(SELECT expense_number FROM expenses WHERE expense_id=%s)) ORDER BY audit_log_id",(str(eid),eid),True)
        assert events==[('CREATE',),('UPDATE',),('APPROVE',),('MARK_PAID',)],events
        wb=load_workbook(BytesIO(c.get('/finance/reports/export?year=2026&status=PAID').data))
        cells=[cell for row in wb['Expense Records'] for cell in row if cell.value=='=1+1']
        assert cells and all(cell.data_type=='s' for cell in cells)
        assert not any(cell.data_type=='f' for ws in wb for row in ws for cell in row)
        assert budgets.get_department_budget_report(1)[0]['spent_amount']==Decimal('0.51')
        assert post('/finance/expenses/add',dict(form,description='Cancelled fixture')).status_code==302
        cancelled=sql('SELECT MAX(expense_id) FROM expenses',fetch=True)[0][0]
        assert post(f'/finance/expenses/{cancelled}/cancel').status_code==302
        assert any(r['expense_id']==cancelled and r['status']=='CANCELLED' for r in reports.report(selected)['expense_records'])
        print('PASS report snapshot during real concurrent commit, full create/edit/approve/pay and cancellation lifecycle, audit integrity and serialized safe XLSX')

        verify_functional_corrections(erp, sql, c, post)

        viewer=client(2)
        assert viewer.get('/finance/reports/export?year=2026').location.endswith('/dashboard')
        sql("INSERT INTO role_permissions SELECT 2,permission_id FROM permissions WHERE permission_code='finance.view'")
        assert viewer.get('/finance/reports/export?year=2026').status_code==200
        html=viewer.get('/finance/budgets?year=2026').get_data(as_text=True)
        assert '/finance/budgets/add' not in html and not re.search(r'/finance/budgets/\d+/edit',html)
        assert viewer.get('/finance/budgets/add').location.endswith('/dashboard')
        before=sql('SELECT COUNT(*) FROM audit_logs',fetch=True)
        assert viewer.post('/finance/budgets/add',data=dict(department_id=1,budget_year=2027,budget_amount='1',csrf_token=csrf_token_for(viewer))).location.endswith('/dashboard')
        assert sql('SELECT COUNT(*) FROM audit_logs',fetch=True)==before
        html=c.get('/finance/budgets?year=2026').get_data(as_text=True)
        assert '/finance/budgets/add' in html and re.search(r'/finance/budgets/\d+/edit',html)
        assert viewer.post(f'/finance/expenses/{cancelled}/cancel',data={'csrf_token':csrf_token_for(viewer)}).location.endswith('/dashboard')
        anon=erp.app.test_client();assert anon.get('/finance/reports/export').location.endswith('/login')
        before=sql('SELECT COUNT(*) FROM audit_logs',fetch=True)
        assert c.post('/finance/expenses/add',data=form).status_code==400
        sql("DELETE rp FROM role_permissions rp JOIN permissions p USING(permission_id) WHERE p.permission_code='finance.view'")
        assert c.get('/finance/reports/export?year=2026').location.endswith('/dashboard')
        assert sql('SELECT COUNT(*) FROM audit_logs',fetch=True)==before
        print('PASS authorized/anonymous/denied/revoked export permissions, response headers and CSRF compatibility')
    print('SUCCESS: Finance reporting verified using disposable MySQL only')


if __name__=='__main__':
    harness.run=run
    harness.main()
