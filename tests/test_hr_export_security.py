"""Synthetic XLSX regressions; no company database connections."""
from collections import defaultdict
from copy import deepcopy
from datetime import date, datetime
from decimal import Decimal
from contextlib import ExitStack
from io import BytesIO
import importlib
import unittest
from unittest.mock import patch

from openpyxl import load_workbook

from leave_export import generate_leave_report_excel
from help_desk_export import generate_helpdesk_report_excel
from security_helpers import stamp_session, user_state


def workbooks(text):
    leave = defaultdict(str, **{key: text for key in (
        'request_number', 'employee_number', 'first_name', 'last_name',
        'department_name', 'leave_name', 'status', 'reason',
        'requested_by_username', 'reviewed_by_username', 'rejection_reason')})
    leave.update(start_date=date(2026, 1, 1), end_date=date(2026, 1, 2),
                 total_days=Decimal('1.5'), created_at=datetime(2026, 1, 1), reviewed_at=None)
    ticket = defaultdict(str, **{key: text for key in (
        'ticket_number', 'employee_name', 'employee_number', 'department_name',
        'category_name', 'subject', 'description', 'priority', 'status',
        'assigned_to_username', 'created_by_username', 'resolution')})
    ticket.update(created_at=datetime(2026, 1, 1), resolved_at=None, closed_at=None)
    category = dict.fromkeys(['total_tickets', 'open_count', 'in_progress_count', 'resolved_count', 'closed_count', 'cancelled_count'], 0)
    category.update(category_code=text, category_name=text)
    technician = {key: value for key, value in category.items() if key not in ('category_code', 'category_name')}
    technician['technician_name'] = text
    leave_type = dict.fromkeys(['total_requests', 'pending_count', 'approved_count', 'rejected_count', 'cancelled_count', 'approved_days'], 0)
    leave_type.update(leave_code=text, leave_name=text)
    before = deepcopy((leave, ticket, category, technician, leave_type))
    streams = [
        generate_leave_report_excel(2026, defaultdict(int), [leave_type],
                                    [leave, deepcopy(leave)], defaultdict(lambda: text)),
        generate_helpdesk_report_excel(defaultdict(int), [category], [technician],
                                       [ticket, deepcopy(ticket)]),
    ]
    assert before == (leave, ticket, category, technician, leave_type), 'Export mutated inputs'
    return [load_workbook(stream, data_only=False) for stream in streams]


class HRExportSecurityTests(unittest.TestCase):
    def test_untrusted_text_is_never_an_xlsx_formula(self):
        for text in ['=1+1', '+SUM(1,2)', '-1+2', '@SUM(A1)', ' \t=1+1',
                     '\r\n=1+1', '\ufeff=1+1', '\u200b=1+1', '\u00a0=1+1']:
            with self.subTest(text=repr(text)):
                for book in workbooks(text):
                    for sheet in book:
                        self.assertFalse(any(c.data_type == 'f' for row in sheet for c in row))
                    sheet, coordinate = ('Leave Records', 'B2') if 'Leave Records' in book.sheetnames else ('Ticket Register', 'C2')
                    self.assertEqual(book[sheet][coordinate].value, text.replace("\r", "\n"))
                    self.assertEqual(book[sheet][coordinate].data_type, 's')

    def test_ordinary_empty_nullable_and_typed_values_survive_repeated_exports(self):
        for text in ['EMP-001', "O'Neil", '日本語', '#N/A', '', None]:
            for _ in range(2):
                leave, desk = workbooks(text)
                self.assertEqual(leave.sheetnames, ['Summary', 'By Leave Type', 'Leave Records'])
                self.assertEqual(desk.sheetnames, ['Summary', 'By Category', 'By Technician', 'Ticket Register'])
                self.assertEqual(leave['Leave Records'].max_row, 3)
                self.assertEqual(desk['Ticket Register'].max_row, 3)
                for cell in [leave['Leave Records']['B2'], desk['Ticket Register']['C2']]:
                    self.assertEqual(cell.value, text or None)
                    self.assertNotEqual(cell.data_type, 'f')
                self.assertEqual(leave['Leave Records']['H2'].value, 1.5)
                self.assertEqual(leave['Leave Records']['H2'].data_type, 'n')
                self.assertEqual(leave['Leave Records']['F2'].value, datetime(2026, 1, 1))
                self.assertEqual(leave['Leave Records']['F2'].number_format, 'yyyy-mm-dd')
                self.assertEqual(desk['Ticket Register']['L2'].value, datetime(2026, 1, 1))

    def test_xml_illegal_controls_are_visible_and_do_not_break_exports(self):
        for text, expected in [('\x00=1+1', '\\x00=1+1'), ('\x0b+1', '\\x0b+1'),
                               ('\ud800=1', '\\ud800=1'), ('\uffff=1', '\\uffff=1')]:
            for book in workbooks(text):
                sheet, coordinate = ('Leave Records', 'B2') if 'Leave Records' in book.sheetnames else ('Ticket Register', 'C2')
                self.assertEqual(book[sheet][coordinate].value, expected)
                self.assertEqual(book[sheet][coordinate].data_type, 's')


class HRExportRouteTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.guard = ExitStack()
        cls.addClassCleanup(cls.guard.close)
        cls.guard.enter_context(patch('dotenv.load_dotenv', return_value=False))
        cls.guard.enter_context(patch.dict('os.environ', SECRET_KEY='synthetic-hr-export-tests', ERP_ENV='testing'))
        cls.guard.enter_context(patch('mysql.connector.connect', side_effect=AssertionError('Company DB forbidden')))
        cls.erp = importlib.import_module('app')

    def setUp(self):
        self.stack = ExitStack()
        self.addCleanup(self.stack.close)
        self.stack.enter_context(patch.dict(self.erp.app.config, TESTING=True))
        self.stack.enter_context(patch("database.ensure_ready"))  # Route unit test: schema integration tested separately.
        self.state = user_state(permissions=['leave.view', 'helpdesk.view'])
        self.stack.enter_context(patch.object(self.erp, 'get_authorization_state', side_effect=lambda uid: self.state))
        self.audit = self.stack.enter_context(patch.object(self.erp, 'create_audit_log'))
        self.client = self.erp.app.test_client()
        with self.client.session_transaction() as session:
            stamp_session(self.erp.app, session, self.state)
        self.reads = {}
        for name, value in {
            'get_leave_report_summary': defaultdict(int), 'get_leave_report_by_type': [],
            'get_leave_report_records': [], 'get_all_employees': [], 'get_all_leave_types': [],
            'get_helpdesk_summary': defaultdict(int), 'get_helpdesk_report_by_category': [],
            'get_helpdesk_report_by_technician': [], 'get_all_helpdesk_tickets': [],
        }.items():
            self.reads[name] = self.stack.enter_context(patch.object(self.erp, name, return_value=value))

    def test_downloads_keep_headers_workbooks_filters_and_audit_behavior(self):
        paths = [('/leave/reports/export?year=2025&status=APPROVED&employee_id=9&leave_type_id=2', 'Leave_Report_2025.xlsx'),
                 ('/it/help-desk/reports/export', 'Help_Desk_Report.xlsx')]
        for path, filename in paths:
            response = self.client.get(path)
            self.addCleanup(response.close)
            self.assertEqual(response.status_code, 200)
            self.assertIn(filename, response.headers['Content-Disposition'])
            self.assertIn('attachment', response.headers['Content-Disposition'])
            self.assertEqual(response.mimetype, 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
            self.assertTrue(load_workbook(BytesIO(response.data)).sheetnames)
        self.reads['get_leave_report_records'].assert_called_once_with(
            year=2025, status='APPROVED', employee_id=9, leave_type_id=2)
        self.audit.assert_not_called()  # Existing export routes do not emit business-action events.

    def test_independent_permissions_revocation_and_disabled_accounts(self):
        for permission, path in [('leave.view', '/leave/reports/export'),
                                 ('helpdesk.view', '/it/help-desk/reports/export')]:
            self.state['permissions'] = frozenset({'helpdesk.view' if permission == 'leave.view' else 'leave.view'})
            self.assertTrue(self.client.get(path).location.endswith('/dashboard'))
        self.state['permissions'] = frozenset()
        for path in ['/leave/reports/export', '/it/help-desk/reports/export']:
            self.assertTrue(self.client.get(path).location.endswith('/dashboard'))
        self.state['is_active'] = False
        self.assertTrue(self.client.get('/leave/reports/export').location.endswith('/login'))
        for read in self.reads.values():
            read.assert_not_called()
        self.audit.assert_not_called()

    def test_csrf_and_method_rejection_preserved(self):
        for path in ['/leave/reports/export', '/it/help-desk/reports/export']:
            self.assertEqual(self.client.post(path).status_code, 405)
        for path in ['/leave/requests/add', '/it/help-desk/add']:
            self.assertEqual(self.client.post(path).status_code, 400)
        for read in self.reads.values():
            read.assert_not_called()
        self.audit.assert_not_called()
