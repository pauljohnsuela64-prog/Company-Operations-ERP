# ==========================================================
# LEAVE MANAGEMENT - EXCEL REPORT EXPORT
# ==========================================================

from io import BytesIO

from excel_text import append_text_row

from openpyxl import Workbook

from openpyxl.styles import (
    Alignment,
    Font,
    PatternFill,
)

from openpyxl.utils import (
    get_column_letter,
)


# ==========================================================
# STYLES
# ==========================================================

TITLE_FILL = PatternFill(
    "solid",
    fgColor="1F4E78",
)

HEADER_FILL = PatternFill(
    "solid",
    fgColor="D9EAF7",
)


# ==========================================================
# STYLE HEADER
# ==========================================================

def _style_header(
    worksheet,
    row_number,
):

    for cell in worksheet[row_number]:

        cell.font = Font(
            bold=True,
        )

        cell.fill = HEADER_FILL

        cell.alignment = Alignment(
            vertical="center",
        )


# ==========================================================
# AUTO SIZE
# ==========================================================

def _auto_size_columns(
    worksheet,
):

    for column_number, column_cells in enumerate(
        worksheet.columns,
        start=1,
    ):

        max_length = 0

        column_letter = (
            get_column_letter(
                column_number
            )
        )


        for cell in column_cells:

            if cell.value is None:
                continue

            max_length = max(
                max_length,
                len(
                    str(cell.value)
                ),
            )


        worksheet.column_dimensions[
            column_letter
        ].width = min(
            max_length + 2,
            45,
        )


# ==========================================================
# GENERATE LEAVE REPORT
# ==========================================================

def generate_leave_report_excel(
    selected_year,
    summary,
    type_summary,
    records,
    filter_labels,
):

    workbook = Workbook()


    # ======================================================
    # SUMMARY
    # ==========================================================

    summary_sheet = workbook.active

    summary_sheet.title = "Summary"


    summary_sheet.merge_cells(
        "A1:B1"
    )

    summary_sheet["A1"] = (
        f"Leave Report - {selected_year}"
    )

    summary_sheet["A1"].font = Font(
        bold=True,
        color="FFFFFF",
        size=14,
    )

    summary_sheet["A1"].fill = (
        TITLE_FILL
    )

    summary_sheet["A1"].alignment = Alignment(
        horizontal="center",
    )


    append_text_row(summary_sheet, [])


    # ======================================================
    # FILTER INFORMATION
    # ==========================================================

    append_text_row(summary_sheet,
        [
            "Filter",
            "Value",
        ]
    )

    append_text_row(summary_sheet,
        [
            "Year",
            selected_year,
        ]
    )

    append_text_row(summary_sheet,
        [
            "Status",
            filter_labels["status"],
        ]
    )

    append_text_row(summary_sheet,
        [
            "Employee",
            filter_labels["employee"],
        ]
    )

    append_text_row(summary_sheet,
        [
            "Leave Type",
            filter_labels["leave_type"],
        ]
    )


    _style_header(
        summary_sheet,
        3,
    )


    append_text_row(summary_sheet, [])

    append_text_row(summary_sheet,
        [
            "Metric",
            "Value",
        ]
    )

    append_text_row(summary_sheet,
        [
            "Total Requests",
            summary["total_requests"],
        ]
    )

    append_text_row(summary_sheet,
        [
            "Pending",
            summary["pending_count"],
        ]
    )

    append_text_row(summary_sheet,
        [
            "Approved",
            summary["approved_count"],
        ]
    )

    append_text_row(summary_sheet,
        [
            "Rejected",
            summary["rejected_count"],
        ]
    )

    append_text_row(summary_sheet,
        [
            "Cancelled",
            summary["cancelled_count"],
        ]
    )

    append_text_row(summary_sheet,
        [
            "Approved Leave Days",
            summary["approved_days"],
        ]
    )


    _style_header(
        summary_sheet,
        9,
    )


    # ======================================================
    # BY LEAVE TYPE
    # ==========================================================

    type_sheet = workbook.create_sheet(
        "By Leave Type"
    )


    append_text_row(type_sheet,
        [
            "Code",
            "Leave Type",
            "Total Requests",
            "Pending",
            "Approved",
            "Rejected",
            "Cancelled",
            "Approved Days",
        ]
    )


    for item in type_summary:

        append_text_row(type_sheet,
            [
                item["leave_code"],
                item["leave_name"],
                item["total_requests"],
                item["pending_count"],
                item["approved_count"],
                item["rejected_count"],
                item["cancelled_count"],
                item["approved_days"],
            ]
        )


    _style_header(
        type_sheet,
        1,
    )


    # ======================================================
    # LEAVE RECORDS
    # ==========================================================

    record_sheet = workbook.create_sheet(
        "Leave Records"
    )


    append_text_row(record_sheet,
        [
            "Request Number",
            "Employee Number",
            "Employee",
            "Department",
            "Leave Type",
            "Start Date",
            "End Date",
            "Total Days",
            "Status",
            "Reason",
            "Requested By",
            "Created At",
            "Reviewed By",
            "Reviewed At",
            "Rejection Reason",
        ]
    )


    for record in records:

        employee_name = (
            f"{record['first_name']} "
            f"{record['last_name']}"
        )


        append_text_row(record_sheet,
            [
                record[
                    "request_number"
                ],

                record[
                    "employee_number"
                ],

                employee_name,

                record[
                    "department_name"
                ] or "",

                record[
                    "leave_name"
                ],

                record[
                    "start_date"
                ],

                record[
                    "end_date"
                ],

                float(
                    record[
                        "total_days"
                    ]
                ),

                record[
                    "status"
                ],

                record[
                    "reason"
                ],

                record[
                    "requested_by_username"
                ],

                record[
                    "created_at"
                ],

                record[
                    "reviewed_by_username"
                ] or "",

                record[
                    "reviewed_at"
                ],

                record[
                    "rejection_reason"
                ] or "",
            ]
        )


    _style_header(
        record_sheet,
        1,
    )


    # ======================================================
    # DATE FORMATS
    # ==========================================================

    for row_number in range(
        2,
        record_sheet.max_row + 1,
    ):

        record_sheet.cell(
            row=row_number,
            column=6,
        ).number_format = (
            "yyyy-mm-dd"
        )

        record_sheet.cell(
            row=row_number,
            column=7,
        ).number_format = (
            "yyyy-mm-dd"
        )

        record_sheet.cell(
            row=row_number,
            column=12,
        ).number_format = (
            "yyyy-mm-dd hh:mm:ss"
        )

        record_sheet.cell(
            row=row_number,
            column=14,
        ).number_format = (
            "yyyy-mm-dd hh:mm:ss"
        )


    # ======================================================
    # WRAP LONG TEXT
    # ==========================================================

    for row_number in range(
        2,
        record_sheet.max_row + 1,
    ):

        record_sheet.cell(
            row=row_number,
            column=10,
        ).alignment = Alignment(
            vertical="top",
            wrap_text=True,
        )

        record_sheet.cell(
            row=row_number,
            column=15,
        ).alignment = Alignment(
            vertical="top",
            wrap_text=True,
        )


    # ======================================================
    # FREEZE HEADERS
    # ==========================================================

    type_sheet.freeze_panes = "A2"

    record_sheet.freeze_panes = "A2"


    # ======================================================
    # EXCEL FILTERS
    # ==========================================================

    type_sheet.auto_filter.ref = (
        type_sheet.dimensions
    )

    record_sheet.auto_filter.ref = (
        record_sheet.dimensions
    )


    # ======================================================
    # AUTO SIZE
    # ==========================================================

    for worksheet in workbook.worksheets:

        _auto_size_columns(
            worksheet
        )


    # ======================================================
    # RETURN FILE
    # ==========================================================

    output = BytesIO()

    workbook.save(
        output
    )

    output.seek(0)

    return output