# ==========================================================
# IT ASSET - EXCEL REPORT EXPORT
# ==========================================================

from io import BytesIO
from decimal import Decimal

from openpyxl.utils import get_column_letter
from openpyxl import Workbook
from openpyxl.styles import (
    Alignment,
    Font,
    PatternFill,
)


# ==========================================================
# STYLES
# ==========================================================

TITLE_FILL = PatternFill(
    "solid",
    fgColor="1F4E78",
)

HEADER_FILL = PatternFill(
    "solid",
    fgColor="D9EAF7",
)


# ==========================================================
# VALUE HELPERS
# ==========================================================

def _number(value):

    if value is None:
        return 0

    if isinstance(value, Decimal):
        return float(value)

    return value


# ==========================================================
# STYLE TABLE HEADER
# ==========================================================

def _style_header(worksheet, row_number):

    for cell in worksheet[row_number]:

        cell.font = Font(
            bold=True,
        )

        cell.fill = HEADER_FILL

        cell.alignment = Alignment(
            vertical="center",
        )


# ==========================================================
# AUTO SIZE COLUMNS
# ==========================================================

def _auto_size_columns(worksheet):

    for column_number, column_cells in enumerate(
        worksheet.columns,
        start=1,
    ):

        max_length = 0

        # ==================================================
        # GET COLUMN LETTER SAFELY
        #
        # Do not rely on column_cells[0].column_letter
        # because merged cells can be MergedCell objects.
        # ==========================================================

        column_letter = get_column_letter(
            column_number
        )


        # ==================================================
        # FIND LONGEST VALUE
        # ==========================================================

        for cell in column_cells:

            if cell.value is None:
                continue

            value_length = len(
                str(cell.value)
            )

            max_length = max(
                max_length,
                value_length,
            )


        # ==================================================
        # SET WIDTH
        # ==========================================================

        worksheet.column_dimensions[
            column_letter
        ].width = min(
            max_length + 2,
            35,
        )


# ==========================================================
# GENERATE IT ASSET REPORT
# ==========================================================

def generate_it_asset_report_excel(
    summary,
    category_summary,
    department_summary,
    assets,
):

    workbook = Workbook()


    # ======================================================
    # SUMMARY SHEET
    # ======================================================

    summary_sheet = workbook.active

    summary_sheet.title = "Summary"

    summary_sheet.merge_cells(
        "A1:C1"
    )

    summary_sheet["A1"] = (
        "IT Asset Report"
    )

    summary_sheet["A1"].font = Font(
        bold=True,
        color="FFFFFF",
        size=14,
    )

    summary_sheet["A1"].fill = TITLE_FILL

    summary_sheet["A1"].alignment = Alignment(
        horizontal="center",
    )


    summary_sheet.append([])

    summary_sheet.append(
        [
            "Status",
            "Asset Count",
        ]
    )

    summary_sheet.append(
        [
            "Total Assets",
            summary["total_assets"],
        ]
    )

    summary_sheet.append(
        [
            "Available",
            summary["available_count"],
        ]
    )

    summary_sheet.append(
        [
            "Assigned",
            summary["assigned_count"],
        ]
    )

    summary_sheet.append(
        [
            "Maintenance",
            summary["maintenance_count"],
        ]
    )

    summary_sheet.append(
        [
            "Retired",
            summary["retired_count"],
        ]
    )

    summary_sheet.append(
        [
            "Lost",
            summary["lost_count"],
        ]
    )

    _style_header(
        summary_sheet,
        3,
    )


    # ======================================================
    # CATEGORY SHEET
    # ======================================================

    category_sheet = workbook.create_sheet(
        "By Category"
    )

    category_sheet.append(
        [
            "Category Code",
            "Category",
            "Total",
            "Available",
            "Assigned",
            "Maintenance",
            "Retired",
            "Lost",
            "Purchase Value",
        ]
    )

    for category in category_summary:

        category_sheet.append(
            [
                category["category_code"],
                category["category_name"],
                category["total_assets"],
                category["available_count"],
                category["assigned_count"],
                category["maintenance_count"],
                category["retired_count"],
                category["lost_count"],
                _number(
                    category[
                        "total_purchase_cost"
                    ]
                ),
            ]
        )

    _style_header(
        category_sheet,
        1,
    )


    # ======================================================
    # DEPARTMENT SHEET
    # ======================================================

    department_sheet = workbook.create_sheet(
        "By Department"
    )

    department_sheet.append(
        [
            "Department",
            "Total",
            "Available",
            "Assigned",
            "Maintenance",
            "Retired",
            "Lost",
            "Purchase Value",
        ]
    )

    for department in department_summary:

        department_sheet.append(
            [
                department[
                    "department_name"
                ],
                department[
                    "total_assets"
                ],
                department[
                    "available_count"
                ],
                department[
                    "assigned_count"
                ],
                department[
                    "maintenance_count"
                ],
                department[
                    "retired_count"
                ],
                department[
                    "lost_count"
                ],
                _number(
                    department[
                        "total_purchase_cost"
                    ]
                ),
            ]
        )

    _style_header(
        department_sheet,
        1,
    )


    # ======================================================
    # COMPLETE ASSET REGISTER
    # ======================================================

    asset_sheet = workbook.create_sheet(
        "Asset Register"
    )

    asset_sheet.append(
        [
            "Asset Tag",
            "Asset Name",
            "Category",
            "Brand",
            "Model",
            "Serial Number",
            "Department",
            "Assigned To",
            "Status",
            "Condition",
            "Purchase Date",
            "Purchase Cost",
            "Warranty Expiry",
            "Notes",
        ]
    )

    for asset in assets:

        asset_sheet.append(
            [
                asset[
                    "asset_tag"
                ],
                asset[
                    "asset_name"
                ],
                asset[
                    "category_name"
                ],
                asset[
                    "brand"
                ] or "",
                asset[
                    "model"
                ] or "",
                asset[
                    "serial_number"
                ] or "",
                asset[
                    "department_name"
                ] or "",
                asset[
                    "assigned_employee_name"
                ] or "",
                asset[
                    "status"
                ],
                asset[
                    "condition_status"
                ],
                asset[
                    "purchase_date"
                ],
                _number(
                    asset[
                        "purchase_cost"
                    ]
                ),
                asset[
                    "warranty_expiry"
                ],
                asset[
                    "notes"
                ] or "",
            ]
        )

    _style_header(
        asset_sheet,
        1,
    )


    # ======================================================
    # CURRENCY FORMAT
    # ======================================================

    for row in range(
        2,
        category_sheet.max_row + 1,
    ):

        category_sheet.cell(
            row=row,
            column=9,
        ).number_format = (
            '₱#,##0.00'
        )


    for row in range(
        2,
        department_sheet.max_row + 1,
    ):

        department_sheet.cell(
            row=row,
            column=8,
        ).number_format = (
            '₱#,##0.00'
        )


    for row in range(
        2,
        asset_sheet.max_row + 1,
    ):

        asset_sheet.cell(
            row=row,
            column=12,
        ).number_format = (
            '₱#,##0.00'
        )


    # ======================================================
    # DATE FORMAT
    # ======================================================

    for row in range(
        2,
        asset_sheet.max_row + 1,
    ):

        asset_sheet.cell(
            row=row,
            column=11,
        ).number_format = (
            "yyyy-mm-dd"
        )

        asset_sheet.cell(
            row=row,
            column=13,
        ).number_format = (
            "yyyy-mm-dd"
        )


    # ======================================================
    # FILTERS / FREEZE HEADERS
    # ======================================================

    category_sheet.freeze_panes = "A2"
    department_sheet.freeze_panes = "A2"
    asset_sheet.freeze_panes = "A2"

    category_sheet.auto_filter.ref = (
        category_sheet.dimensions
    )

    department_sheet.auto_filter.ref = (
        department_sheet.dimensions
    )

    asset_sheet.auto_filter.ref = (
        asset_sheet.dimensions
    )


    # ======================================================
    # AUTO WIDTH
    # ======================================================

    for worksheet in workbook.worksheets:

        _auto_size_columns(
            worksheet
        )


    # ======================================================
    # RETURN FILE
    # ======================================================

    output = BytesIO()

    workbook.save(
        output
    )

    output.seek(0)

    return output