# ==========================================================
# FINANCE - EXCEL EXPORT
# ==========================================================

from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment
from openpyxl.utils import get_column_letter
from openpyxl.cell.cell import Cell, ILLEGAL_CHARACTERS_RE
import re


def safe_text_cell(worksheet, value):
    """Serialize risky database text as an explicit XLSX string, never a formula.

    XML cannot represent some controls: show their escaped spelling rather than
    failing the entire export. Legitimate numeric/date objects pass unchanged.
    """
    if not isinstance(value, str):
        return value
    value = ILLEGAL_CHARACTERS_RE.sub(lambda m: f'\\x{ord(m[0]):02x}', value)
    value = re.sub('[\ud800-\udfff\ufffe\uffff]', lambda m: f'\\u{ord(m[0]):04x}', value)
    if value.lstrip('\ufeff\u200b \t\r\n\v\f').startswith(('=', '+', '-', '@')):
        cell = Cell(worksheet, value=value)
        cell.data_type = 's'
        return cell
    return value


# ==========================================================
# AUTO SIZE COLUMNS
# ==========================================================

def auto_size_columns(worksheet):

    for column_cells in worksheet.columns:

        max_length = 0

        column_letter = get_column_letter(
            column_cells[0].column
        )

        for cell in column_cells:

            if cell.value is not None:

                cell_length = len(
                    str(cell.value)
                )

                if cell_length > max_length:
                    max_length = cell_length

        worksheet.column_dimensions[
            column_letter
        ].width = max_length + 2


# ==========================================================
# GENERATE FINANCE REPORT
# ==========================================================

def generate_finance_report_excel(
    year,
    summary,
    monthly_expenses,
    category_expenses,
    expense_records,
    filter_labels=None,
):

    workbook = Workbook()


    # ======================================================
    # SUMMARY SHEET
    # ======================================================

    summary_sheet = workbook.active
    summary_sheet.title = "Summary"

    summary_sheet.append(
        [
            "Finance Report",
            year,
        ]
    )

    if filter_labels:
        labels = [f"{key}: {filter_labels[key]}" for key in ('status', 'category_id', 'department_id', 'search_text')
                  if filter_labels.get(key) is not None]
        if labels:
            summary_sheet['D1'] = 'Filters'
            summary_sheet['E1'] = safe_text_cell(summary_sheet, '; '.join(labels))

    summary_sheet.append([])

    summary_sheet.append(
        [
            "Status",
            "Records",
            "Amount",
        ]
    )

    summary_sheet.append(
        [
            "Total",
            summary["total_expenses"],
            summary["total_amount"],
        ]
    )

    summary_sheet.append(
        [
            "Draft",
            summary["draft_count"],
            summary["draft_amount"],
        ]
    )

    summary_sheet.append(
        [
            "Approved",
            summary["approved_count"],
            summary["approved_amount"],
        ]
    )

    summary_sheet.append(
        [
            "Paid",
            summary["paid_count"],
            summary["paid_amount"],
        ]
    )

    summary_sheet.append(
        [
            "Cancelled",
            summary["cancelled_count"],
            summary["cancelled_amount"],
        ]
    )


    # ======================================================
    # MONTHLY PAID EXPENSES
    # ======================================================

    monthly_sheet = workbook.create_sheet(
        "Monthly Paid Expenses"
    )

    monthly_sheet.append(
        [
            "Month",
            "Paid Expenses",
            "Total Amount",
        ]
    )

    for month in monthly_expenses:

        monthly_sheet.append(
            [
                month["month_name"],
                month["expense_count"],
                month["total_amount"],
            ]
        )


    # ======================================================
    # CATEGORY BREAKDOWN
    # ======================================================

    category_sheet = workbook.create_sheet(
        "Category Breakdown"
    )

    category_sheet.append(
        [
            "Category Code",
            "Category",
            "Paid Expenses",
            "Total Amount",
        ]
    )

    for category in category_expenses:

        category_sheet.append(
            [
                safe_text_cell(category_sheet, category["category_code"]),
                safe_text_cell(category_sheet, category["category_name"]),
                category["expense_count"],
                category["total_amount"],
            ]
        )


    # ======================================================
    # EXPENSE RECORDS
    # ======================================================

    expense_sheet = workbook.create_sheet(
        "Expense Records"
    )

    expense_sheet.append(
        [
            "Expense Number",
            "Date",
            "Category",
            "Department",
            "Supplier",
            "Purchase Order",
            "Description",
            "Amount",
            "Status",
            "Reference Number",
            "Payment Method",
        ]
    )

    for expense in expense_records:

        expense_sheet.append(
            [
                safe_text_cell(expense_sheet, expense["expense_number"]),
                expense["expense_date"],
                safe_text_cell(expense_sheet, expense["category_name"]),
                safe_text_cell(expense_sheet, expense["department_name"] or ""),
                safe_text_cell(expense_sheet, expense["supplier_name"] or ""),
                safe_text_cell(expense_sheet, expense["po_number"] or ""),
                safe_text_cell(expense_sheet, expense["description"]),
                expense["amount"],
                expense["status"],
                safe_text_cell(expense_sheet, expense["reference_number"] or ""),
                expense["payment_method"] or "",
            ]
        )


    # ======================================================
    # FORMATTING
    # ======================================================

    for worksheet in workbook.worksheets:

        for cell in worksheet[1]:

            cell.font = Font(
                bold=True
            )

            cell.alignment = Alignment(
                vertical="center"
            )

        auto_size_columns(
            worksheet
        )


    # ======================================================
    # CURRENCY FORMAT
    # ======================================================

    for row in range(
        4,
        summary_sheet.max_row + 1,
    ):
        summary_sheet.cell(
            row=row,
            column=3,
        ).number_format = '₱#,##0.00'


    for row in range(
        2,
        monthly_sheet.max_row + 1,
    ):
        monthly_sheet.cell(
            row=row,
            column=3,
        ).number_format = '₱#,##0.00'


    for row in range(
        2,
        category_sheet.max_row + 1,
    ):
        category_sheet.cell(
            row=row,
            column=4,
        ).number_format = '₱#,##0.00'


    for row in range(
        2,
        expense_sheet.max_row + 1,
    ):
        expense_sheet.cell(
            row=row,
            column=8,
        ).number_format = '₱#,##0.00'


    # ======================================================
    # CREATE FILE IN MEMORY
    # ======================================================

    output = BytesIO()

    workbook.save(
        output
    )

    output.seek(0)

    return output
